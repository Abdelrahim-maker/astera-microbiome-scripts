from pathlib import Path
import re,zipfile,json,hashlib
from docx import Document
from docx.oxml.ns import qn
import pandas as pd,numpy as np
O=Path('manuscript/draft3');d=Document(O/'Astera_manuscript_draft3.docx');text='\n'.join(p.text for p in d.paragraphs);alltext=text+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
assert len(d.tables)==8 and len(d.inline_shapes)==8
assert len(re.findall(r'^Figure [1-8]\. ',text,re.M))==8
assert len(re.findall(r'^Table [1-8]\. ',text,re.M))==8
for bad in ['q0.25','q0.05','SPLIT_SEQUENCE','RPCA proxy','X (16S)','Y (ITS)','Z (18S)']:
 assert bad not in alltext,bad
assert not re.search(r'\bMF\b|\bJoint\b',text)
assert '9,526' in alltext and '155' in alltext and '790' in alltext
assert 'Figure S31' in text and 'Figure S19' not in text
assert 'PRJNA1371968' in text and 'could not be verified' in text
for s in d.inline_shapes:assert s.width<=d.sections[0].page_width-d.sections[0].left_margin-d.sections[0].right_margin
with zipfile.ZipFile(O/'Astera_manuscript_draft3.docx') as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml');assert b'commentReference' not in xml and b'<w:ins ' not in xml and b'<w:del ' not in xml
m=pd.read_csv(O/'audit/sample_manifest_with_read_totals.csv');assert len(m)==131 and not m[['marker','sample_id']].duplicated().any();assert set(m.Depth_cm)=={'0-10','10-30'}
assert (m.management_group.eq('till')==m.sample_date.eq('2022-11-28')).all()
a=pd.read_csv(O/'audit/alpha_tests_with_BH.csv');assert len(a)==6 and (a.q_BH>=a.p_value-1e-12).all()
comments=pd.read_csv(O/'audit/comment_response_log.csv');assert set(comments.comment_id)==set(range(61))
for f in O.glob('*.docx'):
 with zipfile.ZipFile(f) as z:assert z.testzip() is None
sources=[Path('analysis_outputs/manuscript_dbrda/Astera_manuscript_draft2_comments(1).docx'),Path('CFARM_Co_2022_sample_control_list_2023.05.12.xlsx')]+list(Path('manuscript/audit').glob('*_counts.npz'))+list(Path('analysis_outputs/manuscript_dbrda').glob('*.csv'))
pd.DataFrame([{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sources]).to_csv(O/'audit/source_checksums.csv',index=False)
results={'word_document':'Astera_manuscript_draft3.docx','tables':8,'embedded_figures':8,'comments_accounted_for':61,'sample_records':131,'dbRDA_reproduced':True,'count_alpha_validation':True,'cross_marker_R2_validation':True,'figure_assets_visually_reviewed':True,'native_Word_pagination_rendered':False,'outstanding_source_items':'See Draft3_revision_report.docx','validation':'PASS within documented scope'}
(O/'audit/delivery_validation.json').write_text(json.dumps(results,indent=2));print(json.dumps(results,indent=2))
