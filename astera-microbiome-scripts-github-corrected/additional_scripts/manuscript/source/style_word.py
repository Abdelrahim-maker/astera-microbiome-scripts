from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
for name in ['manuscript','supplementary_material']:
 p=Path('manuscript')/(name+'.docx');d=Document(p)
 for sec in d.sections:
  sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(.85)
  h=sec.header.paragraphs[0];h.text='Soil microbial composition and management | Working preprint draft'
  for r in h.runs:r.font.size=Pt(8)
  f=sec.footer.paragraphs[0];f.alignment=1;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11 if name=='manuscript' else 10);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
 for p in d.paragraphs:
  if p.text.startswith(('Figure ','Table ')):p.paragraph_format.keep_with_next=True
 for t in d.tables:
  t.autofit=True
  for row in t.rows:
   for c in row.cells:
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(3)
     for r in p.runs:r.font.size=Pt(9)
  trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 d.save(Path('manuscript')/(name+'.docx'))
print('Word styling complete')
