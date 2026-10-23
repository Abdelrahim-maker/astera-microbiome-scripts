from pathlib import Path
import sys,re,json,html,unicodedata,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import pandas as pd
from PIL import Image
O=Path(__file__).resolve().parents[1];T=O/'tables';F=O/'figures';R=json.loads((O/'audit/references_verified.json').read_text())
TITLE='Soil microbial diversity and association networks under tillage and no-tillage in Ohio soybean fields'
AUTH='Afaf Abdelrahim¹, Alec Ogg³, Klaus Lorenz², Rattan Lal², M. Scott Demyan³, Virginia I. Rich¹'
LABEL={'16S':'16S','ITS':'ITS','AMF':'18S (AMF)'}
# Crossref/DOI metadata checked against supplied references; normalize title case to APA sentence case.
TITLES={
'Fierer2017':'Embracing the unknown: Disentangling the complexities of the soil microbiome',
'Hartman2018':'Cropping practices manipulate abundance patterns of root and soil microbiome members paving the way to smart farming',
'Sommermann2018':'Fungal community profiles in agricultural soils of a long-term field trial under different tillage, fertilization and crop rotation conditions analyzed by high-throughput ITS-amplicon sequencing',
'Bowles2017':'Ecological intensification and arbuscular mycorrhizas: A meta-analysis of tillage and cover crop effects',
'Sun2018':'Tillage changes vertical distribution of soil bacterial and fungal communities',
'Callahan2016':'DADA2: High-resolution sample inference from Illumina amplicon data',
'McMurdie2013':'phyloseq: An R package for reproducible interactive analysis and graphics of microbiome census data',
'Quast2013':'The SILVA ribosomal RNA gene database project: Improved data processing and web-based tools',
'Nilsson2019':'The UNITE database for molecular identification of fungi: Handling dark taxa and parallel taxonomic classifications',
'Opik2010':'The online database MaarjAM reveals global and ecosystemic distribution patterns in arbuscular mycorrhizal fungi (Glomeromycota)',
'Caporaso2012':'Ultra-high-throughput microbial community analysis on the Illumina HiSeq and MiSeq platforms',
'Gardes1993':'ITS primers with enhanced specificity for basidiomycetes—Application to the identification of mycorrhizae and rusts',
'Lee2008':'Improved PCR primers for the detection and identification of arbuscular mycorrhizal fungi',
'Simon1992':'Specific amplification of 18S fungal ribosomal genes from vesicular-arbuscular endomycorrhizal fungi colonizing roots',
'White1990':'Amplification and direct sequencing of fungal ribosomal RNA genes for phylogenetics',
'Martin2011':'Cutadapt removes adapter sequences from high-throughput sequencing reads',
'Hurlbert1971':'The nonconcept of species diversity: A critique and alternative parameters',
'Anderson2001':'A new method for non-parametric multivariate analysis of variance',
'Anderson2006':'Distance-based tests for homogeneity of multivariate dispersions',
'Anderson2013':'PERMANOVA, ANOSIM, and the Mantel test in the face of heterogeneous dispersions: What null hypothesis are you testing?',
'Benjamini1995':'Controlling the false discovery rate: A practical and powerful approach to multiple testing',
'Mallick2021':'Multivariable association discovery in population-scale meta-omics studies',
'Gloor2017':'Microbiome datasets are compositional: And this is not optional',
'Blondel2008':'Fast unfolding of communities in large networks',
'Weiss2016':'Correlation detection strategies in microbial data sets vary widely in sensitivity and precision',
'Rottjers2018':'From hairballs to hypotheses—Biological insights from microbial networks'}
# Normalize upper-case deposit names to bibliographic capitalization.
for entry in R.values():
 for author in entry.get('author',[]):
  if author.get('family','').isupper():author['family']=author['family'].title()
# Corrada Bravo is a compound surname in the author byline.
for au in R['Mallick2021']['author']:
 if au.get('family')=='Bravo':au.update(family='Corrada Bravo',given='Hector')
def one(x):return x[0] if isinstance(x,list) else str(x or '')
def year(k):return re.search(r'\d{4}$',k).group()
def clean(s):return ' '.join(html.unescape(re.sub('<[^>]+>','',one(s))).split())
def initials(g):
 return ' '.join('-'.join(''.join(t[0].upper()+'.' for t in re.findall(r'[^\W\d_]+',part)) for part in word.split('-')) for word in g.split())
def aname(a):return a['family']+', '+initials(a.get('given',''))
def authors(k):
 names=[aname(a) for a in R[k]['author']]
 if len(names)>20:return ', '.join(names[:19])+', … '+names[-1]
 if len(names)==1:return names[0]
 return ', '.join(names[:-1])+', & '+names[-1]
def citation(k):
 a=R[k]['author'];s=a[0]['family']
 if len(a)==2:s+=' & '+a[1]['family']
 elif len(a)>2:s+=' et al.'
 return s+', '+year(k)
def sortkey(k):return unicodedata.normalize('NFKD',R[k]['author'][0]['family']).encode('ascii','ignore').decode().lower()+year(k)
def cited(text):return set(re.findall(r'@([A-Za-z]+\d{4})',text))
def render_citations(text):
 def repl(m):return '('+'; '.join(citation(k) for k in sorted(re.findall(r'@([A-Za-z]+\d{4})',m.group()),key=sortkey))+')'
 return re.sub(r'\[@[^\]]+\]',repl,text)
def reftext(k):
 m=R[k];beg=authors(k)+f' ({year(k)}). '+TITLES[k]+'. '
 if k=='White1990':return beg+'In M. A. Innis, D. H. Gelfand, J. J. Sninsky, & T. J. White (Eds.), *PCR protocols: A guide to methods and applications* (pp. 315–322). Academic Press. https://doi.org/'+m['DOI']
 journal=clean(m.get('container-title'));vol=str(m.get('volume',''));issue=str(m.get('issue',''));pages=str(m.get('page',m.get('article-number',''))).replace('-','–')
 if k=='Hartman2018':pages='Article 14'
 if k=='Sun2018':pages='Article 699'
 if k=='Gloor2017':pages='Article 2224'
 if k in ['Mallick2021','McMurdie2013','Sommermann2018'] and pages:pages='Article '+pages
 return beg+f'*{journal}, {vol}*'+('('+issue+')' if issue else '')+(', '+pages if pages else '')+'. https://doi.org/'+m['DOI']
def rich(p,text):
 # Plain editable citation text and italic journal/volume/species spans.
 for i,s in enumerate(text.split('*')):
  r=p.add_run(s);r.italic=bool(i%2)
 return p
def para(doc,text,style=None):return rich(doc.add_paragraph(style=style),render_citations(text))
def setup(short):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
 sec.header_distance=sec.footer_distance=Inches(.35)
 st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(12);st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(6)
 for name,size in [('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.keep_with_next=True
 d.styles['Caption'].font.name='Times New Roman';d.styles['Caption'].font.size=Pt(10);d.styles['Caption'].font.italic=False
 d.styles['Caption'].paragraph_format.line_spacing=1.0
 h=sec.header.paragraphs[0];h.text=short;h.style=d.styles['Caption']
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.append(field)
 d.core_properties.title=TITLE;d.core_properties.author='Afaf Abdelrahim; Alec Ogg; Klaus Lorenz; Rattan Lal; M. Scott Demyan; Virginia I. Rich';d.core_properties.subject='Microbiology research manuscript';d.core_properties.comments=''
 settings=d.settings.element
 no=OxmlElement('w:autoHyphenation');no.set(qn('w:val'),'false');settings.append(no)
 return d

def table(d,title,headers,rows,note,widths=None,font=10):
 p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run(title).bold=True
 t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.8/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ridx,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if ridx==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
   for key,v in [('top','55'),('bottom','55'),('left','65'),('right','65')]:e=OxmlElement('w:'+key);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');marg.append(e)
   tcpr.append(marg)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.size=Pt(font);r.bold=(ridx==0)
   if ridx==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF2');tcpr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['bottom']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CCD1D6');borders.append(el)
   tcpr.append(borders)
 p=para(d,note);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1
 for r in p.runs:r.font.size=Pt(10)
 return t

def figure(d,name,title,caption,maxheight=6.3):
 d.add_page_break();p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run(title).bold=True
 im=Image.open(F/(name+'.png'));w=min(6.8,maxheight*im.width/im.height)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(F/(name+'.png')),width=Inches(w))
 p=para(d,caption,'Caption');p.paragraph_format.keep_together=True

def fmt(v,n=3,zero=False):
 s=f'{float(v):.{n}f}'
 return s if zero else re.sub(r'^0\.', '.',s)
def refs(d,keys):
 d.add_page_break();d.add_heading('References',level=1)
 for k in sorted(keys,key=sortkey):
  p=para(d,reftext(k));p.paragraph_format.left_indent=Inches(.3);p.paragraph_format.first_line_indent=Inches(-.3);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(7)

main_text=(O/'source/manuscript.md').read_text().replace('prespecified .05 adjusted threshold','.05 adjusted threshold')
(O/'source/manuscript.md').write_text(main_text)
keys=cited(main_text);assert keys.issubset(R),sorted(keys-set(R))
main=setup('Soil microbiomes under tillage and no-tillage')
for line in main_text.splitlines():
 if not line:continue
 if line.startswith('# '):main.add_heading(line[2:],0)
 elif line.startswith('### '):main.add_heading(line[4:],2)
 elif line.startswith('## '):main.add_heading(line[3:],1)
 else:para(main,line)
refs(main,keys)
main.add_page_break();main.add_heading('Tables',1)
z=pd.read_csv(T/'Table1_coverage.csv')
table(main,'Table 1. Analytical coverage and retained sequencing counts.',['Marker','NT n','Till n','Observed ASVs','Median reads'],[[r.Marker,r.NT,r.Till,f'{r.ASVs:,}',f'{r.Median_reads:,.1f}'.removesuffix('.0')] for r in z.itertuples()], 'NT, no-tillage. ASVs are features with at least one count in the retained cohort; zero-total columns are excluded. Read totals are summed processed counts per sample, not raw sequencing yields. The marker cohorts comprise 45 distinct identifiers in total (Table S1).',[1.3,.8,.8,1.5,2.4])
main.add_paragraph()
z=pd.read_csv(T/'Table2_alpha.csv')
table(main,'Table 2. Rarefied alpha-diversity comparisons.',['Marker','Metric','NT median','Till median','p','BH q'],[[LABEL[r.domain],'Richness' if r.metric=='Observed' else 'Shannon',f'{r.nt_median:.3f}',f'{r.till_median:.3f}',fmt(r.p_value),fmt(r.q_BH)] for r in z.itertuples()], 'Group medians of sample means over 200 rarefaction draws. Depths were 2,802 reads (16S), 10,488 (ITS), and 1,829 (AMF). Two-sided Mann–Whitney U tests used asymptotic inference with tie and continuity corrections; BH adjustment covers six comparisons. Tests do not account for profile dependence or design covariates.',[1.2,1.1,1.2,1.2,1,1.1])
main.add_page_break()
z=pd.read_csv(T/'Table3_beta_restricted.csv')
table(main,'Table 3. Management-only PERMANOVA with whole-profile permutations.',['Marker','n','R²','Pseudo-F','p','BH q'],[[LABEL[r.marker],r.n,fmt(r.R2,4,True),fmt(r.pseudo_F,3,True),fmt(r.p,4),fmt(r.q_BH,4)] for r in z.itertuples()], 'Bray–Curtis distances of relative abundances; 9,999 whole-profile permutations within site code and identical depth coverage. Model df = 1; residual df = 42, 41, and 42, respectively. BH adjustment covers the three marker tests. No result meets q ≤ .05. Site and depth coverage restrict permutations but are not fitted covariates; collection date is not a predictor or permutation stratum. Restrictions preserve recorded profiles and depth positions but do not adjust away the confounding of management and collection date. Unrestricted tests and dispersion statistics are retained in Table S3.',[1.5,.6,1.0,1.2,1.2,1.3])
figure(main,'Figure1_alpha','Figure 1. Rarefied alpha diversity by management.', 'Richness (A–C) and Shannon entropy (D–F) for 16S, ITS, and AMF-targeted 18S. Each point is a sample mean across 200 draws at the stated depth. Boxes show medians and interquartile ranges; whiskers extend to the most extreme values within 1.5 interquartile ranges. Green denotes NT and orange tillage. Labels show two-sided rank-sum p values and BH q values across six tests. Sample sizes are 33/11, 33/10, and 34/10 (NT/tillage).',5.8)
figure(main,'Figure2_composition','Figure 2. Community composition and restricted management tests.', 'PCoA of relative-abundance Bray–Curtis dissimilarities for 16S (A), ITS (B), and AMF (C). Green circles denote NT and orange triangles tillage. Lingoes correction was applied for visualization; axis percentages use total corrected inertia. R² and BH q values refer to management-only PERMANOVA on the original distances, with 9,999 whole-profile permutations within site code and depth coverage (Table 3), not to tests of the displayed axes. No confidence regions are inferred from the scatter.',4.5)
figure(main,'Figure3_differential_abundance','Figure 3. Differential-abundance discovery counts.', 'Numbers of ASVs associated with management in the original management-only MaAsLin2 models at q ≤ .05 and q ≤ .25. The broader threshold includes the stricter set; bars are not stacked. BH correction was applied within each marker. The respective totals tested were 915, 127, and 87. These models did not include site, depth, date, or profile effects; the counts describe exploratory fitted-model associations. Coefficients and standard errors for the q ≤ .05 features are in Table S4; all tests are in Supplementary Data.',4.5)
figure(main,'Figure4_network_metrics','Figure 4. Descriptive network topology by management.', 'Connected ASV count, mean degree, and density of the original marker-specific networks. Green denotes NT and orange tillage. Each bar summarizes one inferred graph, not independent replicates. Edges require Spearman |ρ| ≥ .6 and BH q < .05 after count-plus-one CLR transformation. Sample numbers and candidate-feature sets differ between groups. Full statistics and balanced-subset summaries are in Tables S5 and S6.',4.5)
main.save(O/'Astera_Manuscript_bioRxiv.docx')
# Supplement: only metadata and checks relevant to the retained analyses.
supp=setup('Supplementary material — soil microbiomes under tillage and no-tillage')
supp.add_heading('Supplementary material',0);supp.add_heading(TITLE,1);para(supp,AUTH)
para(supp,'Corresponding author: Virginia I. Rich; rich.270@osu.edu')
para(supp,'Contents: Supplementary Methods; Tables S1–S7; Figures S1–S2; companion Supplementary Data workbook. Collection dates remain in sample metadata to document the sampling design; no date-effect or soil-carbon analyses are included.')
supp.add_heading('Supplementary Methods',1)
para(supp,'The manuscript focuses on three marker-specific comparisons. The main beta-diversity inference uses the archived whole-profile restricted management test; the original unrestricted tests remain available as exploratory sensitivity results. The restricted tests use seeds 20260930 (16S), 20260932 (ITS), and 20260934 (AMF), with 9,999 permutations. Pseudonymous profile IDs encode recorded depth pairing rather than verified independent farms. The two site codes are not synonymous with management: Co_MT includes some NT samples.')
para(supp,'The alpha-diversity analysis uses sample means across 200 rarefaction draws. BH correction is applied across six marker-by-metric tests separately for each rarefaction scenario. Rarefaction sampling variation is not independent biological replication. Supplementary Data includes sample means, per-draw summaries, and unique-ASV retention summaries. Balanced beta-diversity and network checks use 20 overlapping NT subsets; all tillage samples are retained. These checks are descriptive and do not remove profile dependence or management–date confounding.')
supp.add_heading('Processing provenance and reproducibility scope',2)
para(supp,'Available 16S records include QIIME2 DADA2 paired-end workflows with run-specific truncation lengths of 151/151 and 142/151, alongside an R workflow using 148/148. The archived ITS R workflow uses primer-trimmed reads, maxEE = 2/2, and no active truncLen setting. The archived AMF workflow uses truncLen = 270/255 and trimLeft = 10. These alternatives are retained as provenance records rather than asserted to be one uniform final pipeline. The AMF script computes a chimera-filtered table but assigns taxonomy to another table; its full linkage to the final analysis object is not established. This does not demonstrate that the final object contains chimeras, but it prevents a complete raw-processing provenance claim.')
para(supp,'Exact database releases, the final run-to-object mapping, and the original complete R environment require author verification. The retained-table analysis and figure preparation do not repair those historical gaps. The primary rarefaction and restricted-management results agree with the previously validated processed-input reproductions. Original MaAsLin2 and R network models were not refitted during this editorial revision. The MaAsLin2 revision corrects the model description; it does not introduce an adjusted model. The simplified network presentation retains the original thresholded associations and balanced subsampling, without hub or taxon-association overlays. Tables and figures were curated from the recorded outputs, and complete nonsignificant MaAsLin2 results are retained.')
supp.add_heading('Supplementary Data guide',2)
para(supp,'Supplementary_Data.xlsx contains full-precision tables for the retained analyses. MaAsLin2_all_results contains all 1,129 tested marker-specific associations, including coefficients, standard errors, p values, q values, and supplied taxonomy. S1_sample_membership documents all 45 sample IDs and marker inclusion. Other sheets contain rarefaction summaries, restricted and unrestricted management tests, network metrics, and balanced-subset results. No carbon measurements or date-effect model results are included in this workbook.')
supp.add_page_break()
z=pd.read_csv(T/'S1_sample_membership.csv')
table(supp,'Table S1. Sample identities, recorded grouping, and marker membership.',['Sample ID','Profile','Mgmt.','Site','Depth','Collected','16S','ITS','AMF'],[[r.sample_id,r.profile_id,'NT' if r.management_group=='NT' else 'Till',r.site,r.Depth_cm,r.sample_date,r._7,r.ITS,r.AMF] for r in z.itertuples()], 'Depth is in cm. Yes/No indicates presence/absence in the retained analytical matrix, not a verified laboratory failure. There are 45 distinct IDs, 42 shared IDs, and 24 recorded profile IDs per marker. Dates document the confounded design and are not analyzed as outcomes or predictors in this revision. Profile codes are pseudonyms; site codes do not establish independent farm replication.',[1.65,.5,.45,.65,.6,.9,.45,.45,.45],8)
supp.add_page_break()
z=pd.read_csv(T/'S2_alpha_sensitivity.csv')
table(supp,'Table S2. Higher-depth alpha-diversity sensitivity analysis.',['Marker','Metric','Reads','NT/Till n','NT median','Till median','p','BH q'],[[LABEL[r.domain],'Richness' if r.metric=='Observed' else 'Shannon',f'{r.depth:,}',f'{r.n_nt}/{r.n_till}',f'{r.nt_median:.3f}',f'{r.till_median:.3f}',fmt(r.p_value),fmt(r.q_BH)] for r in z.itertuples()], 'The higher depth retains at least 95% of samples within each marker. Values are group medians of sample means over 200 draws. Two-sided rank-sum tests and BH correction follow the primary analysis. Samples below the selected depth are excluded from this sensitivity scenario; absence of significance does not establish equivalence.',[1.05,.85,.7,.75,.95,.95,.7,.85],9)
supp.add_paragraph()
z=pd.read_csv(T/'S3_unrestricted_beta_dispersion.csv')
table(supp,'Table S3. Original unrestricted management tests and dispersion checks.',['Marker','R²','Pseudo-F','PERMANOVA p','Dispersion F','Dispersion p'],[[LABEL[r.marker],fmt(r.R2,4,True),fmt(r.pseudo_F,3,True),fmt(r.p),fmt(r.dispersion_F,3,True),fmt(r.dispersion_p)] for r in z.itertuples()], 'The original marker-specific PERMANOVA files supply the p values shown here (999 unrestricted permutations). PERMDISP statistics come from the archived sample-balance analysis. Separate unrestricted reruns have Monte Carlo variation in PERMANOVA p values and are not mixed into these columns. Values are nominal, without across-marker correction. These exploratory sample-level tests do not account for the recorded profile structure.',[1.1,.8,1,1.35,1.2,1.35],9)
supp.add_page_break()
z=pd.read_csv(T/'S4_MaAsLin2_q005.csv')
def taxshort(r):
 if r.marker=='AMF':return 'Glomus cf. microaggregatum'
 a=[x.strip() for x in r.taxonomy.split(';')]
 g=[x[3:] for x in a if x.startswith('g__')];s=[x[3:] for x in a if x.startswith('s__')]
 if g:return g[0]+(' '+s[0] if s else '')
 return re.sub(r'^[a-z]__','',a[-1])
table(supp,'Table S4. All management-associated ASVs at q ≤ .05.',['Marker / ASV','Supplied annotation','Coefficient','SE','p','BH q'],[[r.marker+' / '+r.feature,taxshort(r),fmt(r.coef,3,True),fmt(r.stderr,3,True),f'{r.pval:.2g}',f'{r.qval:.4f}'] for r in z.itertuples()], 'Original management-only MaAsLin2 linear models with NT as reference. Positive coefficients indicate higher modeled transformed relative abundance under tillage; negative coefficients indicate higher NT. SE, model standard error. BH correction is within marker (915, 127, and 87 tested features). Taxonomic ranks vary; annotations are not independent functional validation. Full taxonomy and all nonsignificant results are provided in Supplementary Data.',[1.0,2.25,1,.65,.85,1.05],9)
supp.add_page_break()
z=pd.read_csv(T/'S5_network_metrics.csv')
table(supp,'Table S5. Original descriptive network statistics.',['Network','Nodes','Edges','Density','Mean degree'],[[r.label,str(r.nodes),f'{r.edges:,}',fmt(r.density,3,True),fmt(r.mean_degree,2,True)] for r in z.itertuples()], 'AMF denotes the AMF-targeted 18S dataset. Nodes are incident to retained edges; isolated candidates are excluded. Density = 2E/[N(N − 1)] and mean degree = 2E/N. Each row describes one graph, not a replicated treatment mean.',[1.4,1.1,1.1,1.6,1.6],9)
supp.add_paragraph()
z=pd.read_csv(T/'S6_network_balanced.csv')
table(supp,'Table S6. Median network statistics across 20 balanced subsets.',['Network','Nodes','Edges','Density','Mean degree'],[[r.domain+' '+r.label,f'{r.median_nodes:g}',f'{r.median_edges:g}',fmt(r.median_density,3,True),fmt(r.median_mean_degree,2,True)] for r in z.itertuples()], 'Each iteration retains all tillage samples and draws the same number of NT samples without replacement: 11/11 for 16S, 10/10 for ITS, and 10/10 for AMF. Medians are calculated separately for each metric and need not describe a single realized graph. Overlapping subsets are not independent replicates; no node-level p values are used as tests of field-level management effects.',[1.4,1.1,1.1,1.6,1.6],9)
supp.add_paragraph()
z=pd.read_csv(T/'S7_beta_balanced.csv')
table(supp,'Table S7. Balanced-subset beta-diversity sensitivity.',['Marker','n per group','Median R²','Median p','p < .05 iterations'],[[LABEL[r.domain],11 if r.domain=='16S' else 10,fmt(r.median_permanova_r2,4,True),fmt(r.median_permanova_p,4),f'{round(r.frac_permanova_p_lt_0_05*20)}/20'] for r in z.itertuples()], 'Original unrestricted sample-level PERMANOVA across 20 overlapping balanced subsets. Frequencies are descriptive sensitivity summaries, not power estimates or independent confirmations. These analyses are distinct from the whole-profile restricted tests in main Table 3.',[1.3,1.2,1.2,1.2,1.9],9)
figure(supp,'FigureS1_rarefaction','Figure S1. Read-based rarefaction curves.', 'Exact expected ASV richness under sampling without replacement. Top panels show the full retained libraries on logarithmic read axes; bottom panels show the common range up to the primary rarefaction depth. Green denotes NT and orange tillage. Dashed lines indicate 2,802 reads (16S), 10,488 (ITS), and 1,829 (AMF). Curves are conditional on the processed matrices, do not extrapolate, and do not establish complete recovery of soil diversity.',5.8)
figure(supp,'FigureS2_balance','Figure S2. Sensitivity of unrestricted community tests to sample balance.', 'Each point represents one of 20 overlapping balanced subsets. Boxes show medians and interquartile ranges, with 1.5-interquartile-range whiskers. Panels show PERMANOVA R² and nominal p values; the dashed line marks .05. Sample sizes are 11/11 for 16S and 10/10 for ITS and AMF. Detection frequencies are 18/20, 7/20, and 4/20. These are sample-level unrestricted tests, not the profile-restricted main analysis; sampling confounding persists.',4.6)
supp.save(O/'Astera_Supplementary_Material.docx')
# Author-facing change log, separate from submission files.
change=setup('Author revision notes — not part of the manuscript');change.add_heading('Revision summary and submission notes',0)
para(change,'Prepared October 23, 2026. Original manuscripts, scripts, data and public repositories were preserved. The two new Word documents are submission-formatted drafts, with embedded figures and editable tables; they are not a certification that unresolved study provenance has been verified.')
for head,body in [
('Kept','Three marker datasets; rarefied richness and Shannon diversity; marker-specific Bray–Curtis PCoA; management PERMANOVA and dispersion checks; management-only MaAsLin2; marker-specific networks; essential rarefaction and sample-balance checks. Four main figures, three main tables, two supplementary figures and seven supplementary tables.'),
('Removed','All carbon measurements, carbon models, carbon interactions, carbon-network overlays, carbon results and interpretations; date-only and date-inclusive dbRDA; all variance partitioning; cross-marker PCA, combined PERMANOVA and cross-marker networks; redundant original alpha/NMDS figures; unrelated references and unsupported functional/biocontrol claims. The main analysis workbook excludes carbon values and date-effect model outputs. Historical source copies are retained only under audit/.'),
('Retained sampling information','Collection dates and the management–date confounding statement remain, as required to describe this design honestly. Pseudonymous profile IDs and marker membership are supplied without source farm identities. Funding/affiliation names containing Carbon are retained because they are proper institutional names, not carbon analyses.'),
('Corrected statistics and prose','Main beta-diversity table and Figure 2 now emphasize the existing whole-profile restricted management test. No result passes BH q ≤ .05: 16S q = .0555, ITS q = .3770, AMF q = .3018. Original unrestricted p values remain in Table S3. This editorial choice was not a preregistered analysis plan. MaAsLin2 is correctly described as management-only, not adjusted for depth/site/profile. The q ≤ .05 counts remain 5/14/1; all 1,129 tests are preserved. AMF dispersion p is .074. Sample accounting is 45 unique identifiers, 42 shared, and marker totals 44/43/44; unverified 22-site and sequencing-failure statements were removed.'),
('Figure map','Old main Figure 1 → new Figure 1; old Figure 2 → new Figure 2, with cross-marker panel removed and restricted-test labels; old Figure 5 → new Figure 3; old Figure 6 → new Figure 4, without cross-marker networks; old Figure 9 → Figure S1; old Figure 7 → Figure S2, without cross-marker analysis. Network diagrams, hub labels, differential-abundance outlines, transitivity and modularity have been removed. Only node/edge counts, density and mean degree are retained, with the balanced-subsample comparison. All captions and citations were renumbered.'),
('References and authors','Same author order and affiliations as the supplied manuscript. Virginia I. Rich is corresponding author, using rich.270@osu.edu from her official Ohio State profile (https://microbiology.osu.edu/people/rich.270). References use APA author–date citations, sentence-case titles, complete author lists under APA rules, italic journal/volume formatting, hanging indents, and DOI links. Metadata was retrieved through DOI/Crossref records. The Lee et al. AMF-primer reference was checked by title and corrected to DOI 10.1111/j.1574-6941.2008.00531.x.'),
('Author checks before submission','Confirm the final preprocessing lineage, exact primer/library records and database versions with the laboratory records; the current manuscript openly documents the remaining provenance gap. Confirm recorded profile grouping and management terminology against field records. All authors must review the scientific interpretation, affiliations and preserved funding statement. A competing-interest declaration was not supplied or confirmed, so none was invented; provide the accurate declaration during submission and add it to the manuscript if required.'),
('Availability and upload','Submit Astera_Manuscript_bioRxiv.docx as the manuscript, Astera_Supplementary_Material.docx as supplementary material, and Supplementary_Data.xlsx as supplementary data. Do not upload this revision-notes document or the audit/ directory as manuscript supplements. This revision does not claim that the older Zenodo DOI contains the new final-reference outputs. Review the converted submission PDF for page layout and figure readability before completing submission. No public submission or repository update has been performed.'),
('Formatting verification scope','Word files contain embedded figures, editable tables, page-number fields and no tracked edits or comments. The earlier layout previews are historical and do not represent this revision. Document structure and numerical consistency are checked by the validation script. Native Word/LibreOffice pagination was not rendered, so Word pagination is not claimed as visually verified; check the submission-system PDF. The official bioRxiv FAQ endpoint returned HTTP 403 during this preparation, so current file-size/UI details are not asserted.')]:
 change.add_heading(head,1);para(change,body)
change.save(O/'Revision_Summary_and_Author_Checks.docx')
# Human-readable version and references for inspection.
(O/'audit/manuscript_text_APA.txt').write_text(render_citations(main_text)+'\n\nReferences\n\n'+'\n\n'.join(reftext(k) for k in sorted(keys,key=sortkey)))
(O/'audit/reference_check.csv').write_text('key,doi\n'+'\n'.join(k+','+R[k]['DOI'] for k in sorted(keys,key=sortkey))+'\n')
print('Saved manuscript, supplementary material and author revision notes as Word documents.')
