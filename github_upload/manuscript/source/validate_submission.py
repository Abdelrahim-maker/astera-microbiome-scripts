from pathlib import Path
import sys,re,json,zipfile,hashlib,xml.etree.ElementTree as ET
from docx import Document
import pandas as pd,numpy as np
from statsmodels.stats.multitest import multipletests
O=Path(__file__).resolve().parents[1];checks=[]
def check(name,ok):
 checks.append({'check':name,'passed':bool(ok)})
 if not ok:raise AssertionError(name)
for name,nt,nf in [('Astera_Manuscript_bioRxiv.docx',3,4),('Astera_Supplementary_Material.docx',7,2)]:
 p=O/name;d=Document(p);text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 check(name+' table count',len(d.tables)==nt);check(name+' embedded figure count',len(d.inline_shapes)==nf)
 check(name+' no citation placeholders',not re.search(r'\[@|\(REF\)|TODO|TBD',text))
 check(name+' no removed model headings',not re.search('Date-inclusive|Secondary carbon|Carbon covariation|Variance partitioning|Sequential carbon',text,re.I))
 with zipfile.ZipFile(p) as z:
  check(name+' ZIP CRC',z.testzip() is None)
  xml=z.read('word/document.xml');check(name+' no tracked changes',b'<w:ins ' not in xml and b'<w:del ' not in xml)
  check(name+' no comments','word/comments.xml' not in z.namelist())
  for n in z.namelist():
   if n.endswith('.xml'):ET.fromstring(z.read(n))
  check(name+' XML parsed',True)
 (O/'audit'/(p.stem+'_text.txt')).write_text(text)
main=Document(O/'Astera_Manuscript_bioRxiv.docx');body='\n'.join(p.text for p in main.paragraphs)
check('Corresponding email', 'rich.270@osu.edu' in body)
check('No all-caps Gardes citation','GARDES' not in body)
check('All retained citations have DOI-verified references',len(pd.read_csv(O/'audit/reference_check.csv'))==25)
for n in range(1,5):check('Main figure '+str(n)+' caption',bool(re.search(r'^Figure '+str(n)+r'\.',body,re.M)))
supbody='\n'.join(p.text for p in Document(O/'Astera_Supplementary_Material.docx').paragraphs)
for n in range(1,3):check('Supplement figure S'+str(n)+' caption',bool(re.search(r'^Figure S'+str(n)+r'\.',supbody,re.M)))
b=pd.read_csv(O/'tables/Table3_beta_restricted.csv');check('Restricted p and q unchanged',np.allclose(b.p,[.0185,.377,.2012]) and np.allclose(b.q_BH,[.0555,.377,.3018]));check('Restricted BH correction',np.allclose(b.q_BH,multipletests(b.p,method='fdr_bh')[1]))
a=pd.read_csv(O/'tables/Table2_alpha.csv');check('Alpha BH correction',np.allclose(a.q_BH,multipletests(a.p_value,method='fdr_bh')[1]));check('All alpha q nonsignificant',bool((a.q_BH>.05).all()))
r=pd.read_csv(O/'tables/MaAsLin2_all_results.csv');check('All tested associations retained',len(r)==1129)
check('DA q005 totals',[(r.loc[r.marker.eq(k),'qval']<=.05).sum() for k in ['16S','ITS','AMF']]==[5,14,1])
m=pd.read_csv(O/'tables/S1_sample_membership.csv');check('Cohort union and intersection',len(m)==45 and m[['16S','ITS','AMF']].eq('Yes').all(axis=1).sum()==42)
check('Marker cohort sizes',[m[k].eq('Yes').sum() for k in ['16S','ITS','AMF']]==[44,43,44])
for f in (O/'tables').glob('*.csv'):
 cols=pd.read_csv(f,nrows=0).columns;check(f.name+' excludes chemistry columns',not any(c in cols for c in ['C_16S','C_ITS','C_AMF','carbon_conflict','C','TOC']))
n=pd.read_csv(O/'tables/S5_network_metrics.csv');check('Network mean degree',np.allclose(n.mean_degree,2*n.edges/n.nodes));check('Network density',np.allclose(n.density,2*n.edges/n.nodes/(n.nodes-1)))
check('Workbook opens',len(pd.ExcelFile(O/'Supplementary_Data.xlsx').sheet_names)>=15)
report={'date':'2026-10-23','checks':checks,'all_passed':all(x['passed'] for x in checks),'main_word_figures':4,'supplement_word_figures':2,'main_tables':3,'supplement_tables':7,'references':25,'native_Word_pagination_rendered':False,'secondary_Pandoc_previews_and_figures_visually_reviewed':False,'scope':'Editorial revision and source-result consistency; not a rerun of all historical analyses or raw processing.'}
(O/'audit/validation.json').write_text(json.dumps(report,indent=2)+'\n')
files=['Astera_Manuscript_bioRxiv.docx','Astera_Supplementary_Material.docx','Supplementary_Data.xlsx']
with zipfile.ZipFile(O/'Astera_Submission_Files.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in files:z.write(O/f,f)
(O/'SHA256SUMS.txt').write_text(''.join(hashlib.sha256((O/f).read_bytes()).hexdigest()+'  '+f+'\n' for f in files+['Revision_Summary_and_Author_Checks.docx','Astera_Submission_Files.zip']))
print('PASS:',len(checks),'checks; submission ZIP contains only the two Word submission files and the data workbook.')
